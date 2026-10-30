import csv

import pytest
from llama_index.core import Document

from app.ingestion import convert_documents_to_markdown, markdown_blocks, split_documents


def test_conversion_preserves_colliding_stems_and_updates_metadata(settings):
    (settings.RAW_DATA_DIR / "lecture.md").write_text("# Заголовок\nТекст", encoding="utf-8")
    (settings.RAW_DATA_DIR / "lecture.txt").write_text("Другой текст", encoding="utf-8")
    first = convert_documents_to_markdown(settings)
    second = convert_documents_to_markdown(settings)
    assert [d.id_ for d in first] == [d.id_ for d in second]
    assert (settings.PROCESSED_MD_DIR / "lecture.md.md").exists()
    assert (settings.PROCESSED_MD_DIR / "lecture.txt.md").exists()
    with settings.METADATA_CSV_PATH.open() as stream:
        assert len(list(csv.DictReader(stream))) == 2


def test_structural_chunks_keep_heading_and_table_rows():
    text = "# Курс\n## Методы\n" + "Определение метода. " * 120
    text += "\n\n| Метод | Значение |\n| --- | --- |\n"
    text += "\n".join(f"| метод{i} | число{i} |" for i in range(30))
    document = Document(text=text, id_="stable", metadata={"file_name": "lesson.pdf"})
    nodes = split_documents([document], "structural", 64, 8, tokenizer=str.split)
    assert all(len(n.text.split()) <= 48 for n in nodes)
    assert all(n.metadata["file_name"] == "lesson.pdf" for n in nodes)
    tables = [n for n in nodes if n.metadata["is_table"]]
    assert len(tables) > 1
    assert all("Курс/Методы" in n.text and "| Метод | Значение |" in n.text for n in tables)
    assert sum(n.text.count("| метод") for n in tables) == 30
    assert [n.id_ for n in nodes] == [
        n.id_ for n in split_documents([document], "structural", 64, 8, tokenizer=str.split)
    ]


def test_naive_and_structural_differ():
    document = Document(
        text="# First\nAlpha.\n## Second\nBeta.",
        metadata={"file_name": "a.md"},
    )
    naive = split_documents([document], "naive", 128, 8, tokenizer=str.split)
    structural = split_documents([document], "structural", 128, 8, tokenizer=str.split)
    assert len(naive) == 1
    assert len(structural) == 2


def test_oversized_table_fails_instead_of_truncating():
    text = "| Name | Value |\n| --- | --- |\n| name | " + "word " * 80 + "|"
    with pytest.raises(ValueError, match="table row"):
        split_documents(
            [Document(text=text, metadata={"file_name": "a"})],
            "structural",
            64,
            0,
            tokenizer=str.split,
        )


def test_code_fence_is_not_a_table():
    blocks = markdown_blocks("```\n| a | b |\n| --- | --- |\n```\n")
    assert len(blocks) == 1
    assert blocks[0][1] is False


def test_missing_metadata_rejected():
    with pytest.raises(ValueError, match="file_name"):
        split_documents([Document(text="Test")], "naive", 128, 8, tokenizer=str.split)


def test_real_docling_docx_conversion_offline(settings):
    from docx import Document as WordDocument

    word = WordDocument()
    word.add_heading("Линейная алгебра", level=1)
    word.add_paragraph("Матрица состоит из строк и столбцов.")
    word.save(settings.RAW_DATA_DIR / "lecture.docx")
    converted = convert_documents_to_markdown(settings)
    assert len(converted) == 1
    assert "Линейная алгебра" in converted[0].text
    assert "Матрица состоит" in converted[0].text


def test_real_docling_pptx_conversion_offline(settings):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Линейная алгебра"
    slide.placeholders[1].text = "Матрица состоит из строк и столбцов."
    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Примеры матриц"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for cell, value in zip(
        [table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)],
        ["Тип", "Свойство", "Единичная", "Диагональ из единиц"],
        strict=True,
    ):
        cell.text = value
    deck.save(settings.RAW_DATA_DIR / "lecture.PPTX")
    converted = convert_documents_to_markdown(settings)
    assert len(converted) == 1
    document = converted[0]
    assert document.metadata["file_name"] == "lecture.PPTX"
    for expected in (
        "Линейная алгебра",
        "Матрица состоит",
        "Примеры матриц",
        "Единичная",
        "Диагональ из единиц",
    ):
        assert expected in document.text
    nodes = split_documents([document], "structural", 128, 8, tokenizer=str.split)
    assert all(node.metadata["file_name"] == "lecture.PPTX" for node in nodes)
    with settings.METADATA_CSV_PATH.open() as stream:
        assert list(csv.DictReader(stream))[0]["source_format"] == "pptx"
